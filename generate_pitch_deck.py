import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette: Modern Dark Mode & Electric Accents
    C_BG = RGBColor(11, 15, 25)         # Deep Navy / Dark Slate
    C_CARD = RGBColor(22, 30, 49)       # Dark Card
    C_CYAN = RGBColor(0, 212, 255)      # Electric Cyan
    C_PURPLE = RGBColor(168, 85, 247)   # Electric Violet
    C_WHITE = RGBColor(255, 255, 255)
    C_MUTED = RGBColor(156, 163, 175)   # Light Gray
    C_ACCENT = RGBColor(245, 158, 11)   # Warm Gold

    banner_img = r"C:\Users\ASUS\.gemini\antigravity-ide\brain\743c9114-1085-4e36-a9fe-eb171e152960\edtech_promo_banner_1790763507759.jpg"
    kids_img = r"C:\Users\ASUS\.gemini\antigravity-ide\brain\743c9114-1085-4e36-a9fe-eb171e152960\kids_edtech_gamified_1790764974798.jpg"

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=C_CARD, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Title Slide (Hero Pitch)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Left: Pitch copy
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(6.5), Inches(4.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "EVALAI PRO • MULTI-AGENT EDTECH"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CYAN

    p2 = tf1.add_paragraph()
    p2.text = "Corrigez 60 copies en 5 minutes.\nAvec la rigueur d'un jury académique."
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    p2.space_before = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Le premier système multi-agents qui décompose vos énoncés, audite le code sans complaisance, élimine le plagiat et accompagne chaque apprenant de 8 à 88 ans."
    p3.font.size = Pt(15)
    p3.font.color.rgb = C_MUTED
    p3.space_before = Pt(16)

    # Badges
    p4 = tf1.add_paragraph()
    p4.text = "🚀 YC Pitch Deck Edition  •  ⚡ Multi-Agent Consensus  •  🛡️ Anti-Cheat & Plagiat"
    p4.font.size = Pt(12)
    p4.font.color.rgb = C_ACCENT
    p4.space_before = Pt(24)

    # Right: Hero Image
    if os.path.exists(banner_img):
        s1.shapes.add_picture(banner_img, Inches(7.2), Inches(1.2), Inches(5.3), Inches(5.1))

    # ==========================================
    # SLIDE 2: The Core Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)

    tb2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.0))
    tf2 = tb2.text_frame
    p = tf2.paragraphs[0]
    p.text = "LE PROBLÈME : LA CRISE DE L'ÉVALUATION ET DU TEMPS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    p_t = tf2.add_paragraph()
    p_t.text = "Pourquoi les approches actuelles échouent dans l'enseignement"
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

    problems = [
        ("⏳ Le Week-end Volé", "15 à 20 heures par semaine passées à relire du code rébarbatif, créant une fatigue mentale et un feedback en retard."),
        ("🤖 Les LLM Naïfs & Complaisants", "Un prompt unique donne 16/20 à tout le monde, ignore les bugs sournois et hallucine sur les justifications."),
        ("🚨 La Fraude & Copier-Coller", "L'essor de ChatGPT incite au 'hardcoding' et au plagiat de masse. Les enseignants sont désarmés face aux soumissions."),
        ("📉 Le Manque d'Équité", "Le 1er devoir corrigé le samedi matin n'a jamais la même sévérité que le 50ème corrigé le dimanche à 23h.")
    ]

    for i, (title, desc) in enumerate(problems):
        card = add_card(s2, 0.8 + (i * 2.95), 2.2, 2.75, 4.3, border_color=C_CARD)
        tb = s2.shapes.add_textbox(Inches(1.0 + (i * 2.95)), Inches(2.5), Inches(2.35), Inches(3.7))
        tfc = tb.text_frame
        tfc.word_wrap = True
        pt = tfc.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = C_CYAN
        
        pd = tfc.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(13)
        pd.font.color.rgb = C_MUTED
        pd.space_before = Pt(14)

    # ==========================================
    # SLIDE 3: The Solution - MAS (Multi-Agent System)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)

    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.0))
    tf3 = tb3.text_frame
    p = tf3.paragraphs[0]
    p.text = "NOTRE DIFFÉRENCIATION RADICALE : L'ARCHITECTURE MULTI-AGENTS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CYAN

    p_t = tf3.add_paragraph()
    p_t.text = "Un comité d'IA contradictoire au service de la vérité académique"
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

    agents = [
        ("📐 Architecte Pédagogique", "Lit l'énoncé brut (PDF/MD)\nDécompose les notions (Bloom)\nConstruit la grille JSON stricte", C_CYAN),
        ("🔍 Auditeur Technique", "Vérifie l'exactitude du code\nAnalyse la complexité et les perfs\nProduit des preuves irréfutables", C_WHITE),
        ("😈 Challenger Sceptique", "Traque le plagiat et le hardcoding\nTeste les cas limites (edge cases)\nConteste toute note non méritée", C_PURPLE),
        ("⚖️ Juge Arbitre", "Harmonise les avis contradictoires\nFixe la note finale certifiée\nGarantit l'égalité inter-élèves", C_ACCENT),
        ("💡 Mentor Bienveillant", "Traduit le verdict en feedforward\nConseils personnalisés d'apprentissage\nMicro-défis pour progresser", C_CYAN),
    ]

    for i, (title, desc, color) in enumerate(agents):
        card = add_card(s3, 0.8 + (i * 2.38), 2.2, 2.25, 4.4, border_color=color)
        tb = s3.shapes.add_textbox(Inches(0.95 + (i * 2.38)), Inches(2.4), Inches(1.95), Inches(3.9))
        tfc = tb.text_frame
        tfc.word_wrap = True
        pt = tfc.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = color
        
        pd = tfc.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = C_MUTED
        pd.space_before = Pt(12)

    # ==========================================
    # SLIDE 4: Business Model & Packaging Freemium
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)

    tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.0))
    tf4 = tb4.text_frame
    p = tf4.paragraphs[0]
    p.text = "BUSINESS MODEL : DU PRODUCT-LED GROWTH AU CONTRAT B2B"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    p_t = tf4.add_paragraph()
    p_t.text = "Une mécanique virale gratuite, une monétisation naturelle"
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

    plans = [
        ("Plan Gratuit (Hook)", "0 € / mois", [
            "2 devoirs / mois (max 25 copies)",
            "Pipeline mono-agent rapide",
            "Upload d'énoncé PDF simple",
            "Export CSV basique",
            "Objectif : Virilité & adoption solo"
        ], C_MUTED),
        ("Plan Pro Teacher (Cœur)", "29 € à 49 € / mois", [
            "Devoirs & copies illimités",
            "Moteur Multi-Agents (MAS) complet",
            "Détection Plagiat temporelle (qui a copié)",
            "Rapports détaillés & emailing étudiants",
            "Objectif : Le 'must-have' du prof"
        ], C_CYAN),
        ("Plan Campus / Bootcamp", "2 000 € à 10 000 € / an", [
            "Départements entiers & multi-profs",
            "Plagiat cross-promotions (3 ans)",
            "Intégration native LMS (Moodle/Canvas)",
            "Dashboard Directeur des Études",
            "Souveraineté des données & SLA"
        ], C_ACCENT),
    ]

    for i, (p_name, p_price, feats, col) in enumerate(plans):
        card = add_card(s4, 0.8 + (i * 3.95), 2.1, 3.75, 4.6, border_color=col)
        tb = s4.shapes.add_textbox(Inches(1.0 + (i * 3.95)), Inches(2.3), Inches(3.35), Inches(4.2))
        tfc = tb.text_frame
        tfc.word_wrap = True

        pt = tfc.paragraphs[0]
        pt.text = p_name
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = col

        ppr = tfc.add_paragraph()
        ppr.text = p_price
        ppr.font.size = Pt(22)
        ppr.font.bold = True
        ppr.font.color.rgb = C_WHITE
        ppr.space_before = Pt(4)

        for f in feats:
            pf = tfc.add_paragraph()
            pf.text = f"• {f}"
            pf.font.size = Pt(13)
            pf.font.color.rgb = C_MUTED
            pf.space_before = Pt(8)

    # ==========================================
    # SLIDE 5: The Grand Public & Kids 8+ Vision
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)

    tb5 = s5.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.0))
    tf5 = tb5.text_frame
    p = tf5.paragraphs[0]
    p.text = "EXPANSION B2C : L'ÉDUCATION POUR TOUS (EXEMPLE ENFANT DE 8 ANS)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_ACCENT

    p_t = tf5.add_paragraph()
    p_t.text = "Comment le MAS transforme la pensée informatique en jeu d'aventure"
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

    # Left: Explanation
    tb_kids = s5.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.8), Inches(4.8))
    tfk = tb_kids.text_frame
    tfk.word_wrap = True

    pts = [
        ("🤖 Le Mentor Robot Bienveillant", "Parle à l'enfant avec douceur : 'Bravo pour ta fusée ! Peux-tu faire faire deux tours supplémentaires à ton robot ?'"),
        ("👾 Le Monstre 'Edge-Case Max'", "Le rôle du Challenger devient un boss de jeu vidéo : 'J'ai caché un rocher sur ta route, ton code va-t-il l'éviter ?'"),
        ("🧩 Décomposition Conceptuelle en Puzzles", "Pas de syntaxe brute : boucles, conditions et variables deviennent des briques visuelles et des quêtes."),
        ("❤️ La Rétention Parentale (Volonté de Payer)", "Les parents payent 20€ à 50€/mois pour un tuteur IA stimulant qui apprend à réfléchir plutôt qu'à consommer passivement.")
    ]

    for i, (head, text) in enumerate(pts):
        p_h = tfk.add_paragraph() if i > 0 else tfk.paragraphs[0]
        p_h.text = head
        p_h.font.size = Pt(15)
        p_h.font.bold = True
        p_h.font.color.rgb = C_CYAN
        if i > 0: p_h.space_before = Pt(10)

        p_b = tfk.add_paragraph()
        p_b.text = text
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = C_MUTED
        p_b.space_before = Pt(3)

    # Right: Kids image
    if os.path.exists(kids_img):
        s5.shapes.add_picture(kids_img, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.9))

    # ==========================================
    # SLIDE 6: Roadmap & YC North Star
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)

    tb6 = s6.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.0))
    tf6 = tb6.text_frame
    p = tf6.paragraphs[0]
    p.text = "ROADMAP STRATÉGIQUE & NORTH STAR METRIC"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CYAN

    p_t = tf6.add_paragraph()
    p_t.text = "Notre plan d'exécution pour devenir la référence mondiale"
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

    steps = [
        ("Étape 1 : Validation ENIT", "Mois 1", "Tester sur 100+ copies réelles d'ingénieurs à l'ENIT. Valider un taux d'agrément enseignant > 90%."),
        ("Étape 2 : Lancement Freemium", "Mois 2-3", "Mise en ligne de l'outil self-serve pour 50 profs de fac/bootcamps. Mesurer le Weekly Active Teachers."),
        ("Étape 3 : Module CodeQuest Kids", "Mois 4-6", "Déclinaison de l'interface pour les 8-14 ans. Lancement pilote auprès de 100 familles."),
        ("Étape 4 : Déploiement Enterprise", "Mois 6+", "Signature des 3 premières universités et bootcamps tech majeurs en abonnement annuel.")
    ]

    for i, (title, timing, desc) in enumerate(steps):
        card = add_card(s6, 0.8 + (i * 2.95), 2.2, 2.75, 4.3, border_color=C_CARD)
        tb = s6.shapes.add_textbox(Inches(1.0 + (i * 2.95)), Inches(2.4), Inches(2.35), Inches(3.8))
        tfc = tb.text_frame
        tfc.word_wrap = True

        pt = tfc.paragraphs[0]
        pt.text = timing
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = C_ACCENT

        ph = tfc.add_paragraph()
        ph.text = title
        ph.font.size = Pt(17)
        ph.font.bold = True
        ph.font.color.rgb = C_WHITE
        ph.space_before = Pt(6)

        pd = tfc.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = C_MUTED
        pd.space_before = Pt(14)

    # Save presentation
    output_path = r"e:\Enseignement\ENIT\edtech-eval-platform\edtech_eval_pitch_deck.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_deck()
